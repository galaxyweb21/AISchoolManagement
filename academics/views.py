from core.pagination import paginate_queryset
# academics/views.py
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.views.decorators.http import require_POST
from django.http import HttpResponse, JsonResponse
from django.views.decorators.http import require_http_methods
import json
from datetime import datetime, timedelta
from django.core.exceptions import ValidationError
from django.db import IntegrityError
from django.urls import reverse
from django.db import transaction
from django.utils import timezone
from django.core.paginator import Paginator
from django.db.models import Q

from school.models import AcademicTerm, AcademicYear
from .models import *
from .services import *
from .tasks import generate_timetable_task
from students.models import *

from .forms import TimeSlotForm, GESScheduleForm
from .timetable_configuration_forms import TimetableConfigurationForm, TimetableScheduleBlockFormSet
from .timetable_configuration_model import TimetableConfiguration
from academics.services.timetable_configuration import (
    get_or_create_configuration,
    apply_timetable_configuration,
    build_configured_schedule_preview,
    build_schedule_preview_with_blocks,
    get_explicit_blocks,
    save_explicit_blocks,
)
from .models import PromotionRule, PromotionBatch, StudentPromotion, SchoolClass
from academics.services.promotion_service import PromotionService
from academics.services.timetable_readiness import TimetableReadiness
from academics.services.timetable_export import build_timetable_export_context
from academics.services.ges_schedule import (
    generate_ges_standard_timeslots,
    build_ges_schedule_preview,
    summarize_school_day,
    GESScheduleError,
    GES_STANDARD_DEFAULTS,
    GES_TEMPLATE_LABEL,
)


# ============================================================
# AI TIMETABLER LOGIC (Keep as is)
# ============================================================

@login_required
def timetable_workspace(request):
    school = request.user.school
    if not school:
        messages.error(request, "No school associated with your account.")
        return redirect('dashboard')

    active_term = AcademicTerm.objects.filter(
        academic_year__school=school, academic_year__is_active=True, is_active=True
    ).first()

    timetables = Timetable.objects.filter(school=school)
    if active_term:
        timetables = timetables.filter(academic_term=active_term)
    timetables = timetables.order_by('-generated_at')[:10]

    readiness = TimetableReadiness(school, active_term).run()

    context = {
        'active_term': active_term,
        'timetables': timetables,
        'can_generate': request.user.role in ['SUPER_ADMIN', 'SCHOOL_ADMIN'],
        'readiness': readiness,
    }
    return render(request, 'academics/timetable_workspace.html', context)


@login_required
@require_POST
def generate_timetable(request):
    if request.user.role not in ['SUPER_ADMIN', 'SCHOOL_ADMIN']:
        messages.error(request, "You don't have permission to generate a timetable.")
        return redirect('academics:timetable_workspace')

    school = request.user.school
    active_term = AcademicTerm.objects.filter(
        academic_year__school=school, academic_year__is_active=True, is_active=True
    ).first()

    if not active_term:
        messages.error(request, "No active academic term is configured. Set one up in School Settings first.")
        return redirect('academics:timetable_workspace')

    readiness = TimetableReadiness(school, active_term).run()
    if not readiness['ready']:
        messages.error(
            request,
            "Timetable is not ready for generation. Fix the items shown in the Readiness Check first."
        )
        return redirect('academics:timetable_workspace')

    # The AI timetabler can't schedule anything without a weekly period
    # grid. Rather than fail the run outright, seed the GES-standard
    # schedule automatically the first time a school generates with none
    # configured -- non-destructive (only fills in periods that don't
    # already exist), so this is a complete no-op for any school that
    # already has a schedule set up.
    if not TimeSlot.objects.filter(school=school, is_active=True).exists():
        try:
            created_count, _ = generate_ges_standard_timeslots(school)
        except GESScheduleError as e:
            messages.error(request, f"Could not generate a timetable: no time slots are configured, "
                                     f"and the default GES schedule couldn't be generated automatically ({e}). "
                                     f"Set one up from Academics > Time Slots first.")
            return redirect('academics:timetable_workspace')

        if created_count:
            messages.info(
                request,
                f"No time slots were configured yet, so the GES standard schedule "
                f"({created_count} periods) was generated automatically. You can adjust "
                f"it anytime from Academics > Time Slots."
            )
        else:
            # Every (day, period) already existed but all as inactive --
            # generation left them alone rather than resurrecting them.
            messages.error(
                request,
                "No active time slots are configured for this school. "
                "Check Academics > Time Slots -- your existing periods may be marked inactive."
            )
            return redirect('academics:timeslot_list')

    timetable = AITimetableService.create_pending(
        school=school, academic_term=active_term, generated_by=request.user
    )

    try:
        generate_timetable_task.delay(str(timetable.id))
        messages.info(request, "Timetable generation started. This page will update automatically.")
    except Exception:
        messages.warning(request, "Background worker unavailable - generating inline instead.")
        AITimetableService.run(timetable)

    return redirect('academics:timetable_detail', timetable_id=timetable.id)


@login_required
def timetable_export_pdf(request, timetable_id):
    """Download a server-generated timetable PDF using xhtml2pdf."""
    school = request.user.school
    timetable = get_object_or_404(Timetable, id=timetable_id, school=school)

    if timetable.status in ('PENDING', 'RUNNING'):
        messages.warning(request, "The timetable is still being generated. Please try again when it is complete.")
        return redirect('academics:timetable_detail', timetable_id=timetable_id)

    try:
        from io import BytesIO
        from xhtml2pdf import pisa
        from django.template.loader import render_to_string

        context = build_timetable_export_context(
            timetable,
            class_id=request.GET.get('class_id'),
        )
        html = render_to_string('academics/timetable_export_pdf.html', context, request=request)
        result = BytesIO()
        pdf_status = pisa.CreatePDF(html, dest=result, encoding='UTF-8')
        if pdf_status.err:
            return JsonResponse(
                {'error': 'The timetable PDF could not be generated.'},
                status=500,
            )

        filename = context['export_filename_base'] + '.pdf'
        response = HttpResponse(result.getvalue(), content_type='application/pdf')
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response
    except ImportError:
        return JsonResponse(
            {'error': 'xhtml2pdf is not installed. Install the project requirements and try again.'},
            status=500,
        )
    except Exception as exc:
        return JsonResponse(
            {'error': f'Unable to generate the timetable PDF: {exc}'},
            status=500,
        )


@login_required
def timetable_export_word(request, timetable_id):
    """Download an editable timetable Word document using python-docx."""
    school = request.user.school
    timetable = get_object_or_404(Timetable, id=timetable_id, school=school)

    if timetable.status in ('PENDING', 'RUNNING'):
        messages.warning(request, "The timetable is still being generated. Please try again when it is complete.")
        return redirect('academics:timetable_detail', timetable_id=timetable_id)

    try:
        from io import BytesIO
        from django.http import HttpResponse
        from docx import Document
        from docx.enum.section import WD_ORIENT
        from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Inches, Pt

        context = build_timetable_export_context(
            timetable,
            class_id=request.GET.get('class_id'),
        )

        document = Document()
        section = document.sections[0]
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = section.page_height, section.page_width
        section.top_margin = Inches(0.45)
        section.bottom_margin = Inches(0.45)
        section.left_margin = Inches(0.35)
        section.right_margin = Inches(0.35)

        title = document.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = title.add_run(context['school_name'])
        run.bold = True
        run.font.size = Pt(15)

        subtitle = document.add_paragraph()
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = subtitle.add_run(f"{context['timetable_title']} — {context['term_name']}")
        run.bold = True
        run.font.size = Pt(11)

        meta = document.add_paragraph()
        meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
        meta.add_run(
            f"Status: {context['status_label']}   |   Fitness: {context['fitness']}   |   "
            f"Hard Conflicts: {context['hard_conflicts']}   |   Soft Conflicts: {context['soft_conflicts']}   |   "
            f"Entries: {context['entries_count']}"
        ).font.size = Pt(8)

        table = document.add_table(rows=1, cols=6)
        table.style = 'Table Grid'
        table.autofit = True
        headers = ['Period / Time'] + context['days_display']
        for i, header in enumerate(headers):
            cell = table.rows[0].cells[i]
            cell.text = header
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in paragraph.runs:
                    r.bold = True
                    r.font.size = Pt(8)

        for row in context['rows']:
            cells = table.add_row().cells
            if row['type'] == 'block':
                merged = cells[0]
                for cell in cells[1:]:
                    merged = merged.merge(cell)
                block = row['block']
                text = f"{block['label']}  |  {block['start']} – {block['end']}"
                merged.text = text
                merged.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for paragraph in merged.paragraphs:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for r in paragraph.runs:
                        r.bold = True
                        r.font.size = Pt(8)
                continue

            period_label = row['period_label']
            cells[0].text = period_label
            for paragraph in cells[0].paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in paragraph.runs:
                    r.bold = True
                    r.font.size = Pt(7.5)

            for idx, cell_data in enumerate(row['cells'], start=1):
                cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
                if not cell_data:
                    cells[idx].text = ''
                    continue
                lines = [cell_data['subject']]
                if cell_data.get('class'):
                    lines.append(cell_data['class'])
                if cell_data.get('teacher'):
                    lines.append(cell_data['teacher'])
                if cell_data.get('room'):
                    lines.append(cell_data['room'])
                cells[idx].text = '\n'.join(lines)
                for paragraph in cells[idx].paragraphs:
                    for r in paragraph.runs:
                        r.font.size = Pt(7)

        footer = document.add_paragraph()
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.add_run('Generated by EduAI School Management').font.size = Pt(7)

        output = BytesIO()
        document.save(output)
        output.seek(0)
        filename = context['export_filename_base'] + '.docx'
        response = HttpResponse(
            output.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response
    except ImportError:
        return JsonResponse(
            {'error': 'python-docx is not installed. Install the project requirements and try again.'},
            status=500,
        )
    except Exception as exc:
        return JsonResponse(
            {'error': f'Unable to generate the timetable Word document: {exc}'},
            status=500,
        )


@login_required
def timetable_detail(request, timetable_id):
    school = request.user.school
    timetable = get_object_or_404(Timetable, id=timetable_id, school=school)

    if timetable.status in ('PENDING', 'RUNNING'):
        return render(request, 'academics/timetable_generating.html', {'timetable': timetable})

    entries_count = TimetableEntry.objects.filter(timetable=timetable).count()
    entries = TimetableEntry.objects.filter(
        timetable=timetable
    ).select_related('school_class', 'subject', 'teacher__user', 'room', 'timeslot')

    days = ['MON', 'TUE', 'WED', 'THU', 'FRI']
    timeslots = list(TimeSlot.objects.filter(school=school, is_active=True).order_by('period_index', 'day'))
    period_indexes = sorted(set(s.period_index for s in timeslots))
    classes = SchoolClass.objects.filter(school=school).order_by('name')
    selected_class_id = request.GET.get('class_id')
    selected_class_display = None
    if selected_class_id:
        selected_class = SchoolClass.objects.filter(id=selected_class_id, school=school).first()
        if selected_class:
            selected_class_display = selected_class.name

    slot_lookup = {}
    for entry in entries:
        if selected_class_id and str(entry.school_class_id) != selected_class_id:
            continue
        slot_lookup[(entry.timeslot.period_index, entry.timeslot.day)] = entry

    config = get_or_create_configuration(school)
    explicit_blocks = get_explicit_blocks(config)
    # Do not let Save Configuration Only retroactively relabel an existing
    # timetable whose TimeSlots have not been updated. Explicit blocks are
    # shown only when the saved configuration matches the current teaching
    # period grid.
    try:
        expected_slots = build_configured_schedule_preview(config)
        actual_slots = {(s.day, s.period_index, s.start_time, s.end_time) for s in timeslots}
        expected = {(r['day'], r['period_index'], r['start_time'], r['end_time']) for r in expected_slots}
        if actual_slots != expected:
            explicit_blocks = []
    except Exception:
        explicit_blocks = []
    blocks_by_period = {}
    for block in explicit_blocks:
        blocks_by_period.setdefault(block.after_period, []).append(block)

    slot_by_period_day = {(slot.period_index, slot.day): slot for slot in timeslots}

    def block_for_day(after_period, day):
        matches = [b for b in blocks_by_period.get(after_period, []) if b.day in (None, '', day)]
        return matches[0] if matches else None

    def block_cells(after_period):
        cells = []
        for day in days:
            block = block_for_day(after_period, day)
            if not block:
                cells.append(None)
                continue
            if after_period == 0:
                first_slot = slot_by_period_day.get((period_indexes[0], day)) if period_indexes else None
                if not first_slot:
                    cells.append(None)
                    continue
                end = first_slot.start_time
                anchor = datetime.combine(datetime.today(), end) - timedelta(minutes=block.minutes)
                cells.append({'block': block, 'start': anchor.time(), 'end': end, 'minutes': block.minutes})
            else:
                current_slot = slot_by_period_day.get((after_period, day))
                next_slot = slot_by_period_day.get((after_period + 1, day))
                if current_slot and next_slot:
                    cells.append({'block': block, 'start': current_slot.end_time, 'end': next_slot.start_time, 'minutes': block.minutes})
                elif current_slot:
                    anchor = datetime.combine(datetime.today(), current_slot.end_time) + timedelta(minutes=block.minutes)
                    cells.append({'block': block, 'start': current_slot.end_time, 'end': anchor.time(), 'minutes': block.minutes})
                else:
                    cells.append(None)
        return cells

    rows = []
    if 0 in blocks_by_period:
        rows.append({'type': 'block', 'cells': block_cells(0)})

    for index, period in enumerate(period_indexes):
        rows.append({
            'type': 'period',
            'period': period,
            'cells': [slot_lookup.get((period, d)) for d in days],
            'period_slots': [slot_by_period_day.get((period, d)) for d in days],
        })
        if period in blocks_by_period:
            rows.append({'type': 'block', 'cells': block_cells(period)})

    context = {
        'timetable': timetable,
        'days': days,
        'rows': rows,
        'classes': classes,
        'selected_class_id': selected_class_id,
        'selected_class_display': selected_class_display,
        'entries_count': entries_count,
        'can_publish': request.user.role in ['SUPER_ADMIN', 'SCHOOL_ADMIN'],
        'schedule_config': config,
        'schedule_blocks': explicit_blocks,
    }
    return render(request, 'academics/timetable_detail.html', context)


@login_required
@require_http_methods(["GET", "POST"])
def timetable_configuration(request):
    """Configure the school's weekly timetable and explicit non-teaching blocks."""
    school = request.user.school
    if not school:
        return redirect('dashboard')
    if request.user.role not in ['SUPER_ADMIN', 'SCHOOL_ADMIN']:
        messages.error(request, "You don't have permission to change timetable configuration.")
        return redirect('academics:timetable_workspace')

    config = get_or_create_configuration(school)
    if request.method == 'POST':
        form = TimetableConfigurationForm(request.POST, instance=config, config=config)
        block_formset = TimetableScheduleBlockFormSet(
            request.POST, instance=config, form_kwargs={'periods_per_day': int(request.POST.get('periods_per_day') or config.periods_per_day)}
        )
        if form.is_valid() and block_formset.is_valid():
            config = form.save()
            try:
                save_explicit_blocks(config, block_formset.forms)
            except ValueError as exc:
                messages.error(request, str(exc))
                return redirect('academics:timetable_configuration')

            apply_schedule = request.POST.get('apply_schedule') == '1'
            replace_existing = request.POST.get('replace_existing') == '1'
            if apply_schedule:
                try:
                    created, skipped = apply_timetable_configuration(school, config, replace_existing=replace_existing)
                    if replace_existing:
                        messages.success(request, f"Schedule configuration saved and {created} time slots regenerated safely.")
                    elif created:
                        messages.success(request, f"Schedule configuration saved. {created} missing teaching periods were added; existing periods were left untouched.")
                    else:
                        messages.success(request, "Schedule configuration saved. Existing teaching periods were left untouched.")
                except ValueError as exc:
                    messages.error(request, str(exc))
            else:
                messages.success(request, "Timetable configuration saved. Existing TimeSlots were not changed.")
            return redirect('academics:timetable_configuration')
    else:
        form = TimetableConfigurationForm(instance=config, config=config)
        block_formset = TimetableScheduleBlockFormSet(instance=config, form_kwargs={'periods_per_day': config.periods_per_day})

    preview = build_schedule_preview_with_blocks(config)
    preview_by_day = {}
    for row in preview:
        preview_by_day.setdefault(row['day'], []).append(row)

    return render(request, 'academics/timetable_configuration.html', {
        'form': form,
        'block_formset': block_formset,
        'config': config,
        'preview_by_day': preview_by_day,
        'can_manage': True,
        'has_timetable_entries': TimetableEntry.objects.filter(timetable__school=school).exists(),
    })


@login_required
@require_POST
def publish_timetable(request, timetable_id):
    if request.user.role not in ['SUPER_ADMIN', 'SCHOOL_ADMIN']:
        messages.error(request, "You don't have permission to publish a timetable.")
        return redirect('academics:timetable_detail', timetable_id=timetable_id)

    school = request.user.school
    timetable = get_object_or_404(Timetable, id=timetable_id, school=school)
    if timetable.status != 'COMPLETE':
        messages.error(request, "Only a completed timetable can be published.")
        return redirect('academics:timetable_detail', timetable_id=timetable_id)

    AITimetableService.publish(timetable)
    messages.success(request, "Timetable published. It's now the active timetable for the term.")
    return redirect('academics:timetable_detail', timetable_id=timetable.id)


# academics/views.py - Fixed delete_timetable view

@login_required
def delete_timetable(request, timetable_id):
    if request.user.role not in ['SUPER_ADMIN', 'SCHOOL_ADMIN']:
        return JsonResponse({'success': False, 'error': "You don't have permission to delete a timetable."}, status=403)

    school = request.user.school
    timetable = get_object_or_404(Timetable, id=timetable_id, school=school)

    if request.method == 'GET':
        # Count entries directly using the timetable filter
        entries_count = TimetableEntry.objects.filter(timetable=timetable).count()
        return render(request, 'academics/timetable_delete_modal.html', {
            'timetable': timetable,
            'entries_count': entries_count,
            'action_url': 'academics:delete_timetable'
        })

    if timetable.is_published:
        return JsonResponse({
            'success': False,
            'error': "This is the published, active timetable — publish a different one before deleting it.",
        })

    # TimetableEntry rows cascade-delete with the timetable (on_delete=CASCADE),
    # so nothing extra to clean up here.
    timetable.delete()
    return JsonResponse({
        'success': True,
        'message': "Timetable deleted.",
        'redirect_url': reverse('academics:timetable_workspace'),
    })


# ============================================================
# SUBJECT MANAGEMENT - FULL MODAL SUPPORT
# ============================================================

@login_required
def subject_list(request):
    school = request.user.school
    if not school: return redirect('dashboard')
    subjects = Subject.objects.filter(school=school).order_by('name')
    return render(request, 'academics/subject_list.html',
                  {'subjects': paginate_queryset(subjects, request), 'can_manage': request.user.role in ['SUPER_ADMIN', 'SCHOOL_ADMIN']})


@login_required
def subject_create(request):
    if request.user.role not in ['SUPER_ADMIN', 'SCHOOL_ADMIN']:
        return JsonResponse({'success': False, 'error': "Permission denied."}, status=403)
    school = request.user.school
    if request.method == 'GET':
        return render(request, 'academics/subject_form_modal.html',
                      {'mode': 'create', 'action_url': 'academics:subject_create'})
    name = request.POST.get('name', '').strip()
    requires_lab = request.POST.get('requires_lab', False) == 'on'
    if not name:
        return JsonResponse({'success': False, 'error': "Subject name is required."})
    if Subject.objects.filter(school=school, name=name).exists():
        return JsonResponse({'success': False, 'error': f"A subject named '{name}' already exists."})
    Subject.objects.create(school=school, name=name, requires_lab=requires_lab)
    return JsonResponse({'success': True, 'message': f"Subject '{name}' created successfully."})


@login_required
def subject_edit(request, subject_id):
    if request.user.role not in ['SUPER_ADMIN', 'SCHOOL_ADMIN']:
        return JsonResponse({'success': False, 'error': "Permission denied."}, status=403)
    school = request.user.school
    subject = get_object_or_404(Subject, id=subject_id, school=school)
    if request.method == 'GET':
        return render(request, 'academics/subject_form_modal.html',
                      {'mode': 'edit', 'subject': subject, 'action_url': 'academics:subject_edit'})
    name = request.POST.get('name', '').strip()
    requires_lab = request.POST.get('requires_lab', False) == 'on'
    if not name:
        return JsonResponse({'success': False, 'error': "Subject name is required."})
    if Subject.objects.filter(school=school, name=name).exclude(id=subject.id).exists():
        return JsonResponse({'success': False, 'error': f"A subject named '{name}' already exists."})
    subject.name = name
    subject.requires_lab = requires_lab
    subject.save()
    return JsonResponse({'success': True, 'message': f"Subject '{name}' updated successfully."})


@login_required
def subject_delete(request, subject_id):
    if request.user.role not in ['SUPER_ADMIN', 'SCHOOL_ADMIN']:
        return JsonResponse({'success': False, 'error': "Permission denied."}, status=403)
    school = request.user.school
    subject = get_object_or_404(Subject, id=subject_id, school=school)
    if request.method == 'GET':
        return render(request, 'academics/subject_delete_modal.html',
                      {'subject': subject, 'action_url': 'academics:subject_delete'})
    subject.delete()
    return JsonResponse({'success': True, 'message': f"Subject '{subject.name}' deleted successfully."})


# ============================================================
# SCHOOL CLASS MANAGEMENT - FIXED WITH COMPLETE DUPLICATE CHECK
# ============================================================

@login_required
def school_class_list(request):
    school = request.user.school
    if not school: return redirect('dashboard:dashboard')
    classes = SchoolClass.objects.filter(school=school).select_related('grade_level').order_by('grade_level__order',
                                                                                               'name')
    return render(request, 'academics/class_list.html',
                  {'classes': paginate_queryset(classes, request), 'can_manage': request.user.role in ['SUPER_ADMIN', 'SCHOOL_ADMIN']})


@login_required
def school_class_create(request):
    if request.user.role not in ['SUPER_ADMIN', 'SCHOOL_ADMIN']:
        return JsonResponse({'success': False, 'error': "Permission denied."}, status=403)

    school = request.user.school
    grade_levels = GradeLevel.objects.filter(school=school).order_by('order')

    if request.method == 'GET':
        return render(request, 'academics/class_form_modal.html', {
            'grade_levels': grade_levels,
            'mode': 'create',
            'action_url': 'academics:school_class_create'
        })

    # Handle POST request
    try:
        name = request.POST.get('name', '').strip()
        grade_level_id = request.POST.get('grade_level', '').strip()

        if not all([name, grade_level_id]):
            return JsonResponse({
                'success': False,
                'error': "Class name and Grade Level are required."
            }, status=400)

        grade_level = get_object_or_404(GradeLevel, id=grade_level_id, school=school)

        # Check if class exists with same name AND grade level.
        if SchoolClass.objects.filter(school=school, name__iexact=name, grade_level=grade_level).exists():
            return JsonResponse({
                'success': False,
                'error': f"A class named '{name}' already exists in {grade_level.name}. Each class name must be unique within a grade level."
            }, status=400)

        # Create the class
        SchoolClass.objects.create(school=school, name=name, grade_level=grade_level)

        return JsonResponse({
            'success': True,
            'message': f"Class '{name}' created successfully in {grade_level.name}."
        })

    except IntegrityError as e:
        return JsonResponse({
            'success': False,
            'error': (
                f"'{name}' couldn't be saved because of a database-level uniqueness "
                f"constraint that appears to ignore Grade Level. If you're trying to "
                f"reuse this class name in a different grade level, the SchoolClass "
                f"model's unique constraint likely needs to include grade_level "
                f"(this requires a models.py change + migration, not just a view fix)."
            )
        }, status=400)
    except Exception as e:
        return JsonResponse({
            'success': False,
            'error': f"An error occurred: {str(e)}"
        }, status=500)


@login_required
def school_class_edit(request, class_id):
    if request.user.role not in ['SUPER_ADMIN', 'SCHOOL_ADMIN']:
        return JsonResponse({'success': False, 'error': "Permission denied."}, status=403)

    school = request.user.school
    school_class = get_object_or_404(SchoolClass, id=class_id, school=school)
    grade_levels = GradeLevel.objects.filter(school=school).order_by('order')

    if request.method == 'GET':
        return render(request, 'academics/class_form_modal.html', {
            'mode': 'edit',
            'school_class': school_class,
            'grade_levels': grade_levels,
            'action_url': 'academics:school_class_edit'
        })

    # Handle POST request
    try:
        name = request.POST.get('name', '').strip()
        grade_level_id = request.POST.get('grade_level', '').strip()

        if not all([name, grade_level_id]):
            return JsonResponse({
                'success': False,
                'error': "Class name and Grade Level are required."
            }, status=400)

        grade_level = get_object_or_404(GradeLevel, id=grade_level_id, school=school)

        if SchoolClass.objects.filter(school=school, name__iexact=name, grade_level=grade_level).exclude(
                id=school_class.id).exists():
            return JsonResponse({
                'success': False,
                'error': f"A class named '{name}' already exists in {grade_level.name}. Each class name must be unique within a grade level."
            }, status=400)

        school_class.name = name
        school_class.grade_level = grade_level
        school_class.save()

        return JsonResponse({
            'success': True,
            'message': f"Class '{name}' updated successfully in {grade_level.name}."
        })

    except IntegrityError as e:
        return JsonResponse({
            'success': False,
            'error': (
                f"'{name}' couldn't be saved because of a database-level uniqueness "
                f"constraint that appears to ignore Grade Level. If you're trying to "
                f"reuse this class name in a different grade level, the SchoolClass "
                f"model's unique constraint likely needs to include grade_level "
                f"(this requires a models.py change + migration, not just a view fix)."
            )
        }, status=400)
    except Exception as e:
        return JsonResponse({
            'success': False,
            'error': f"An error occurred: {str(e)}"
        }, status=500)


@login_required
def school_class_delete(request, class_id):
    if request.user.role not in ['SUPER_ADMIN', 'SCHOOL_ADMIN']:
        return JsonResponse({'success': False, 'error': "Permission denied."}, status=403)

    school = request.user.school
    school_class = get_object_or_404(SchoolClass, id=class_id, school=school)

    if request.method == 'GET':
        return render(request, 'academics/class_delete_modal.html', {
            'school_class': school_class,
            'action_url': 'academics:school_class_delete'
        })

    # Handle POST request
    try:
        if school_class.timetable_entries.exists():
            return JsonResponse({
                'success': False,
                'error': f"Cannot delete '{school_class.name}' because it has timetable entries. Remove the timetable entries first."
            }, status=400)

        class_name = school_class.name
        school_class.delete()

        return JsonResponse({
            'success': True,
            'message': f"Class '{class_name}' deleted successfully."
        })

    except IntegrityError as e:
        return JsonResponse({
            'success': False,
            'error': f"Database error: {str(e)}"
        }, status=400)
    except Exception as e:
        return JsonResponse({
            'success': False,
            'error': f"An error occurred: {str(e)}"
        }, status=500)


# ============================================================
# ROOM MANAGEMENT
# ============================================================

@login_required
def room_list(request):
    school = request.user.school
    if not school: return redirect('dashboard:dashboard')
    rooms = Room.objects.filter(school=school).order_by('name')
    return render(request, 'academics/room_list.html',
                  {'rooms': paginate_queryset(rooms, request), 'can_manage': request.user.role in ['SUPER_ADMIN', 'SCHOOL_ADMIN']})


@login_required
def room_create(request):
    if request.user.role not in ['SUPER_ADMIN', 'SCHOOL_ADMIN']:
        return JsonResponse({'success': False, 'error': "Permission denied."}, status=403)
    school = request.user.school
    if request.method == 'GET':
        return render(request, 'academics/room_form_modal.html',
                      {'mode': 'create', 'action_url': 'academics:room_create'})
    name = request.POST.get('name', '').strip()
    capacity = request.POST.get('capacity', 40)
    is_lab = request.POST.get('is_lab', False) == 'on'
    if not name:
        return JsonResponse({'success': False, 'error': "Room name is required."})
    if Room.objects.filter(school=school, name=name).exists():
        return JsonResponse({'success': False, 'error': f"A room named '{name}' already exists."})
    Room.objects.create(school=school, name=name, capacity=capacity, is_lab=is_lab)
    return JsonResponse({'success': True, 'message': f"Room '{name}' created successfully."})


@login_required
def room_edit(request, room_id):
    if request.user.role not in ['SUPER_ADMIN', 'SCHOOL_ADMIN']:
        return JsonResponse({'success': False, 'error': "Permission denied."}, status=403)
    school = request.user.school
    room = get_object_or_404(Room, id=room_id, school=school)
    if request.method == 'GET':
        return render(request, 'academics/room_form_modal.html',
                      {'mode': 'edit', 'room': room, 'action_url': 'academics:room_edit'})
    name = request.POST.get('name', '').strip()
    capacity = request.POST.get('capacity', 40)
    is_lab = request.POST.get('is_lab', False) == 'on'
    if not name:
        return JsonResponse({'success': False, 'error': "Room name is required."})
    if Room.objects.filter(school=school, name=name).exclude(id=room.id).exists():
        return JsonResponse({'success': False, 'error': f"A room named '{name}' already exists."})
    room.name = name
    room.capacity = capacity
    room.is_lab = is_lab
    room.save()
    return JsonResponse({'success': True, 'message': f"Room '{name}' updated successfully."})


@login_required
def room_delete(request, room_id):
    if request.user.role not in ['SUPER_ADMIN', 'SCHOOL_ADMIN']:
        return JsonResponse({'success': False, 'error': "Permission denied."}, status=403)
    school = request.user.school
    room = get_object_or_404(Room, id=room_id, school=school)
    if request.method == 'GET':
        return render(request, 'academics/room_delete_modal.html',
                      {'room': room, 'action_url': 'academics:room_delete'})
    room.delete()
    return JsonResponse({'success': True, 'message': f"Room '{room.name}' deleted successfully."})


# ============================================================
# TIMESLOT MANAGEMENT - FULL MODAL SUPPORT
# ============================================================

@login_required
def timeslot_list(request):
    school = request.user.school
    if not school: return redirect('dashboard')
    timeslots = TimeSlot.objects.filter(school=school).order_by('day', 'period_index')
    return render(request, 'academics/timeslot_list.html',
                  {'timeslots': paginate_queryset(timeslots, request), 'can_manage': request.user.role in ['SUPER_ADMIN', 'SCHOOL_ADMIN']})


@login_required
@require_http_methods(["GET", "POST"])
def timeslot_create(request):
    """Create a new timeslot"""
    school = request.user.school

    if request.method == "POST":
        form = TimeSlotForm(request.POST)
        if form.is_valid():
            try:
                timeslot = form.save(commit=False)
                timeslot.school = school
                timeslot.save()

                if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
                    return JsonResponse({
                        'success': True,
                        'timeslot_id': str(timeslot.id),
                        'message': 'TimeSlot created successfully!'
                    })
                return redirect('academics:timeslot_list')
            except IntegrityError:
                if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
                    return JsonResponse({
                        'success': False,
                        'error': 'A timeslot with this day and period already exists.'
                    }, status=400)
                form.add_error(None, 'A timeslot with this day and period already exists.')
        else:
            if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
                return JsonResponse({
                    'success': False,
                    'error': ' '.join([' '.join(errors) for errors in form.errors.values()])
                }, status=400)
    else:
        form = TimeSlotForm()

    return render(request, 'academics/timeslot_form_modal.html', {
        'mode': 'create',
        'action_url': 'academics:timeslot_create',
        'form': form,
        'timeslot': None
    })


@login_required
@require_http_methods(["GET", "POST"])
def timeslot_edit(request, timeslot_id):
    """Edit an existing timeslot"""
    school = request.user.school
    timeslot = get_object_or_404(TimeSlot, id=timeslot_id, school=school)

    if request.method == "POST":
        form = TimeSlotForm(request.POST, instance=timeslot)
        if form.is_valid():
            try:
                form.save()

                if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
                    return JsonResponse({
                        'success': True,
                        'message': 'TimeSlot updated successfully!'
                    })
                return redirect('academics:timeslot_list')
            except IntegrityError:
                if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
                    return JsonResponse({
                        'success': False,
                        'error': 'A timeslot with this day and period already exists.'
                    }, status=400)
                form.add_error(None, 'A timeslot with this day and period already exists.')
        else:
            if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
                return JsonResponse({
                    'success': False,
                    'error': ' '.join([' '.join(errors) for errors in form.errors.values()])
                }, status=400)
    else:
        form = TimeSlotForm(instance=timeslot)

    return render(request, 'academics/timeslot_form_modal.html', {
        'mode': 'edit',
        'action_url': 'academics:timeslot_edit',
        'form': form,
        'timeslot': timeslot
    })


@login_required
@require_http_methods(["GET", "POST"])
def timeslot_delete(request, timeslot_id):
    """Delete a timeslot"""
    school = request.user.school
    timeslot = get_object_or_404(TimeSlot, id=timeslot_id, school=school)

    if request.method == "POST":
        try:
            timeslot.delete()

            if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
                return JsonResponse({
                    'success': True,
                    'message': 'TimeSlot deleted successfully!'
                })
            return redirect('academics:timeslot_list')
        except Exception as e:
            if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
                return JsonResponse({
                    'success': False,
                    'error': f'Error deleting timeslot: {str(e)}'
                }, status=500)
            return render(request, 'academics/timeslot_delete_modal.html', {
                'timeslot': timeslot,
                'action_url': 'academics:timeslot_delete',
                'error': f'Error deleting timeslot: {str(e)}'
            })

    return render(request, 'academics/timeslot_delete_modal.html', {
        'timeslot': timeslot,
        'action_url': 'academics:timeslot_delete'
    })


@login_required
@require_http_methods(["GET", "POST"])
def timeslot_generate_ges(request):
    """
    Bulk-generate a week of TimeSlots from the Ghana Education Service's
    proposed periods allocation (8 x 50-minute periods/day, with two
    daily breaks folded into the timing) -- or a customized variant of
    it. Non-destructive by default: existing (day, period) TimeSlots are
    left alone unless "replace_existing" is checked.
    """
    school = request.user.school
    if not school:
        return redirect('dashboard')

    if request.user.role not in ['SUPER_ADMIN', 'SCHOOL_ADMIN']:
        if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            return JsonResponse({'success': False, 'error': 'You do not have permission to do this.'}, status=403)
        messages.error(request, "You do not have permission to do this.")
        return redirect('academics:timeslot_list')

    if request.method == "POST":
        form = GESScheduleForm(request.POST)
        if form.is_valid():
            try:
                created, skipped = generate_ges_standard_timeslots(
                    school, **form.to_service_kwargs()
                )
            except GESScheduleError as e:
                if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
                    return JsonResponse({'success': False, 'error': str(e)}, status=400)
                form.add_error(None, str(e))
            else:
                if created and skipped:
                    message = f"Generated {created} new time slot(s); {skipped} already existed and were left as-is."
                elif created:
                    message = f"Generated {created} new time slot(s) from the GES standard schedule."
                else:
                    message = "No new time slots were created -- every period already exists. Check 'replace and regenerate' to overwrite them."

                if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
                    return JsonResponse({'success': True, 'message': message})
                messages.success(request, message)
                return redirect('academics:timeslot_list')
        else:
            if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
                return JsonResponse({
                    'success': False,
                    'error': ' '.join([' '.join(errors) for errors in form.errors.values()])
                }, status=400)
    else:
        form = GESScheduleForm()

    preview_error = None
    preview_rows = []
    day_summary = []
    try:
        defaults_kwargs = {
            'start_time': GES_STANDARD_DEFAULTS['start_time'],
            'period_length_minutes': GES_STANDARD_DEFAULTS['period_length_minutes'],
            'periods_per_day': GES_STANDARD_DEFAULTS['periods_per_day'],
            'breaks': GES_STANDARD_DEFAULTS['breaks'],
            'days': ['MON'],
        }
        preview_rows = build_ges_schedule_preview(**defaults_kwargs)
        day_summary = summarize_school_day(
            period_length_minutes=GES_STANDARD_DEFAULTS['period_length_minutes'],
            periods_per_day=GES_STANDARD_DEFAULTS['periods_per_day'],
            breaks=GES_STANDARD_DEFAULTS['breaks'],
        )
    except GESScheduleError as e:
        preview_error = str(e)

    return render(request, 'academics/ges_schedule_modal.html', {
        'form': form,
        'action_url': 'academics:timeslot_generate_ges',
        'template_label': GES_TEMPLATE_LABEL,
        'preview_rows': preview_rows,
        'preview_error': preview_error,
        'day_summary': day_summary,
    })


# ============================================================
# STUDENT PROMOTION VIEWS - FIXED
# ============================================================

@login_required
def promotion_dashboard(request):
    """Dashboard for student promotions."""
    if request.user.role not in ['SUPER_ADMIN', 'SCHOOL_ADMIN', 'HOD']:
        messages.error(request, "You don't have permission to access promotions.")
        return redirect('dashboard')

    school = request.user.school

    # Get current academic year
    current_year = AcademicYear.objects.filter(school=school, is_active=True).first()

    # Get promotion statistics
    stats = PromotionService.get_promotion_statistics(school, current_year)

    # Get recent promotions - FIXED: filter through academic_year__school
    recent_promotions = StudentPromotion.objects.filter(
        promotion_batch__academic_year__school=school
    ).select_related('student__user', 'from_grade_level', 'to_grade_level', 'promotion_batch')[:20]

    # Get pending promotions - FIXED: filter through academic_year__school
    pending = StudentPromotion.objects.filter(
        promotion_batch__academic_year__school=school,
        status='PENDING'
    ).count()

    # Get promotion rules
    rules = PromotionRule.objects.filter(school=school, is_active=True)

    context = {
        'stats': stats,
        'recent_promotions': recent_promotions,
        'pending_count': pending,
        'rules': rules,
        'current_year': current_year,
        'active_tab': 'academics',
    }
    return render(request, 'academics/promotion/dashboard.html', context)


@login_required
def promotion_batch_list(request):
    """List all promotion batches."""
    if request.user.role not in ['SUPER_ADMIN', 'SCHOOL_ADMIN', 'HOD']:
        messages.error(request, "You don't have permission to view promotions.")
        return redirect('dashboard')

    school = request.user.school

    # FIXED: Filter through academic_year__school
    batches = PromotionBatch.objects.filter(
        academic_year__school=school
    ).order_by('-created_at')

    page_obj = paginate_queryset(batches, request)

    context = {
        'batches': page_obj,
        'active_tab': 'academics',
    }
    return render(request, 'academics/promotion/batch_list.html', context)


@login_required
def promotion_batch_create(request):
    """Create and process a new promotion batch."""
    if request.user.role not in ['SUPER_ADMIN', 'SCHOOL_ADMIN']:
        return JsonResponse({'success': False, 'error': "Permission denied."}, status=403)

    school = request.user.school

    if request.method == 'GET':
        # Get available grade levels and terms
        grade_levels = GradeLevel.objects.filter(school=school).order_by('order')
        # FIXED: Filter terms through academic_year__school
        terms = AcademicTerm.objects.filter(
            academic_year__school=school,
            is_active=True
        ).select_related('academic_year')

        rules = PromotionRule.objects.filter(school=school, is_active=True)

        context = {
            'grade_levels': grade_levels,
            'terms': terms,
            'rules': rules,
            'mode': 'create',
            'action_url': 'academics:promotion_batch_create',
        }
        return render(request, 'academics/promotion/batch_form_modal.html', context)

    # POST - Process promotion batch
    from_grade_level_id = request.POST.get('from_grade_level')
    to_grade_level_id = request.POST.get('to_grade_level')
    academic_term_id = request.POST.get('academic_term')
    promotion_rule_id = request.POST.get('promotion_rule')
    batch_name = request.POST.get('name', '').strip()
    mode = request.POST.get('mode', 'AUTO')

    if not all([from_grade_level_id, to_grade_level_id, academic_term_id]):
        return JsonResponse({'success': False, 'error': "All required fields must be filled."})

    try:
        result = PromotionService.process_promotion_batch(
            school=school,
            from_grade_level_id=from_grade_level_id,
            to_grade_level_id=to_grade_level_id,
            academic_term_id=academic_term_id,
            promotion_rule_id=promotion_rule_id or None,
            batch_name=batch_name,
            processed_by=request.user,
            mode=mode
        )

        if result['success']:
            messages.success(
                request,
                f"Promotion batch processed: {result['stats']['promoted']} promoted, "
                f"{result['stats']['conditional']} conditional, "
                f"{result['stats']['repeated']} repeated."
            )
            return redirect('academics:promotion_batch_detail', batch_id=result['batch'].id)
        else:
            messages.error(request, f"Error processing promotion: {result['error']}")
            return redirect('academics:promotion_dashboard')

    except Exception as e:
        messages.error(request, f"Error: {str(e)}")
        return redirect('academics:promotion_dashboard')


@login_required
def promotion_batch_detail(request, batch_id):
    """View details of a promotion batch."""
    if request.user.role not in ['SUPER_ADMIN', 'SCHOOL_ADMIN', 'HOD']:
        messages.error(request, "You don't have permission to view promotions.")
        return redirect('dashboard')

    school = request.user.school

    # FIXED: Filter through academic_year__school
    batch = get_object_or_404(PromotionBatch, id=batch_id, academic_year__school=school)

    # Get promotions in this batch
    promotions = StudentPromotion.objects.filter(
        promotion_batch=batch
    ).select_related('student__user', 'from_grade_level', 'to_grade_level')

    # Backfill missing promotion-result notifications for existing records.
    PromotionService.notify_promotion_results(promotions)

    # Apply filters
    status_filter = request.GET.get('status')
    if status_filter:
        promotions = promotions.filter(status=status_filter)

    page_obj = paginate_queryset(promotions, request)

    context = {
        'batch': batch,
        'promotions': page_obj,
        'selected_status': status_filter,
        'active_tab': 'academics',
    }
    return render(request, 'academics/promotion/batch_detail.html', context)


@login_required
def promotion_apply(request, promotion_id):
    """Apply a promotion to a student."""
    if request.user.role not in ['SUPER_ADMIN', 'SCHOOL_ADMIN']:
        return JsonResponse({'success': False, 'error': "Permission denied."}, status=403)

    try:
        result = PromotionService.apply_promotion(promotion_id, request.user)

        if result['success']:
            messages.success(request, result['message'])
        else:
            messages.error(request, result['error'])

    except Exception as e:
        messages.error(request, f"Error: {str(e)}")

    # Redirect back to the referring page
    referer = request.META.get('HTTP_REFERER', 'academics:promotion_dashboard')
    return redirect(referer)


@login_required
def promotion_bulk_apply(request, batch_id):
    """Apply all eligible promotions; GET returns a confirmation fragment."""
    if request.user.role not in ['SUPER_ADMIN', 'SCHOOL_ADMIN']:
        if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            return JsonResponse({'success': False, 'error': "Permission denied."}, status=403)
        messages.error(request, "You don't have permission to perform this action.")
        return redirect('academics:promotion_dashboard')

    school = request.user.school
    batch = get_object_or_404(PromotionBatch, id=batch_id, academic_year__school=school)

    if request.method == 'POST':
        try:
            promotions = StudentPromotion.objects.filter(
                promotion_batch=batch, status__in=['PROMOTED', 'CONDITIONAL']
            )
            applied = 0
            for promotion in promotions:
                result = PromotionService.apply_promotion(promotion.id, request.user)
                if result['success']:
                    applied += 1
            message = f"Successfully applied {applied} promotions." if applied else "No promotions were applied."
            if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
                return JsonResponse({'success': True, 'message': message, 'redirect_url': reverse('academics:promotion_batch_detail', args=[batch.id])})
            (messages.success if applied else messages.warning)(request, message)
        except Exception as e:
            logger.exception('Bulk promotion application failed for batch %s', batch.id)
            if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
                return JsonResponse({'success': False, 'error': str(e)}, status=400)
            messages.error(request, f"Error: {str(e)}")
        return redirect('academics:promotion_batch_detail', batch_id=batch.id)

    promotions = StudentPromotion.objects.filter(
        promotion_batch=batch, status__in=['PROMOTED', 'CONDITIONAL']
    ).select_related('student__user', 'from_grade_level', 'to_grade_level')

    context = {
        'batch': batch,
        'promotions': promotions,
        'action_url': 'academics:promotion_bulk_apply',
        'active_tab': 'academics',
    }

    # AJAX requests receive only the modal fragment. Direct browser visits
    # receive a complete styled page, so Apply All never exposes raw HTML.
    if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
        return render(request, 'academics/promotion/bulk_apply_modal.html', context)
    return render(request, 'academics/promotion/bulk_apply_page.html', context)


@login_required
def promotion_result_detail(request, promotion_id):
    """Read-only promotion result for its linked parent/student or school administrators."""
    promotion = get_object_or_404(
        StudentPromotion.objects.select_related(
            'student__user', 'student__parent', 'from_grade_level', 'from_school_class',
            'to_grade_level', 'to_school_class', 'promotion_batch', 'approved_by'
        ), id=promotion_id
    )
    user = request.user
    if not (user.role in ['SUPER_ADMIN', 'SCHOOL_ADMIN', 'HOD'] or
            promotion.student.parent_id == user.id or promotion.student.user_id == user.id):
        messages.error(request, "You don't have permission to view this promotion result.")
        return redirect('dashboard')
    PromotionService.notify_promotion_result(promotion)
    return render(request, 'academics/promotion/result_detail.html', {
        'promotion': promotion, 'student': promotion.student, 'active_tab': 'academics'
    })


@login_required
def promotion_rule_list(request):
    """List all promotion rules."""
    if request.user.role not in ['SUPER_ADMIN', 'SCHOOL_ADMIN']:
        messages.error(request, "You don't have permission to view promotion rules.")
        return redirect('dashboard')

    school = request.user.school

    rules = PromotionRule.objects.filter(school=school).order_by('from_grade_level__order')

    context = {
        'rules': paginate_queryset(rules, request),
        'active_tab': 'academics',
    }
    return render(request, 'academics/promotion/rule_list.html', context)


@login_required
def promotion_rule_create(request):
    """Create a new promotion rule."""
    if request.user.role not in ['SUPER_ADMIN', 'SCHOOL_ADMIN']:
        return JsonResponse({'success': False, 'error': "Permission denied."}, status=403)

    school = request.user.school

    if request.method == 'GET':
        grade_levels = GradeLevel.objects.filter(school=school).order_by('order')
        return render(request, 'academics/promotion/rule_form_modal.html', {
            'grade_levels': grade_levels,
            'mode': 'create',
            'action_url': 'academics:promotion_rule_create'
        })

    # POST - Create rule
    from_grade_id = request.POST.get('from_grade_level')
    to_grade_id = request.POST.get('to_grade_level')
    promotion_mode = request.POST.get('promotion_mode', 'AUTO')
    min_pass = request.POST.get('minimum_passing_grade', 50)
    min_subjects = request.POST.get('minimum_subjects_to_pass', 0)
    min_average = request.POST.get('minimum_overall_average', 50)
    min_attendance = request.POST.get('minimum_attendance_percentage', 75)
    allow_conditional = request.POST.get('allow_conditional_promotion') == 'on'
    max_conditional = request.POST.get('max_conditional_subjects', 2)
    eval_term = request.POST.get('evaluation_term_sequence', 3)

    if not all([from_grade_id, to_grade_id]):
        return JsonResponse({'success': False, 'error': "From and To grade levels are required."})

    try:
        rule = PromotionRule.objects.create(
            school=school,
            from_grade_level_id=from_grade_id,
            to_grade_level_id=to_grade_id,
            promotion_mode=promotion_mode,
            minimum_passing_grade=min_pass,
            minimum_subjects_to_pass=min_subjects,
            minimum_overall_average=min_average,
            minimum_attendance_percentage=min_attendance,
            allow_conditional_promotion=allow_conditional,
            max_conditional_subjects=max_conditional,
            evaluation_term_sequence=eval_term,
        )
        return JsonResponse({
            'success': True,
            'message': f"Promotion rule created for {rule.from_grade_level.name} → {rule.to_grade_level.name}"
        })
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)})


@login_required
def promotion_rule_edit(request, rule_id):
    """Edit an existing promotion rule."""
    if request.user.role not in ['SUPER_ADMIN', 'SCHOOL_ADMIN']:
        return JsonResponse({'success': False, 'error': "Permission denied."}, status=403)

    school = request.user.school
    rule = get_object_or_404(PromotionRule, id=rule_id, school=school)

    if request.method == 'GET':
        grade_levels = GradeLevel.objects.filter(school=school).order_by('order')
        return render(request, 'academics/promotion/rule_form_modal.html', {
            'rule': rule,
            'grade_levels': grade_levels,
            'mode': 'edit',
            'action_url': 'academics:promotion_rule_edit'
        })

    # POST - Update rule
    promotion_mode = request.POST.get('promotion_mode', 'AUTO')
    min_pass = request.POST.get('minimum_passing_grade', 50)
    min_subjects = request.POST.get('minimum_subjects_to_pass', 0)
    min_average = request.POST.get('minimum_overall_average', 50)
    min_attendance = request.POST.get('minimum_attendance_percentage', 75)
    allow_conditional = request.POST.get('allow_conditional_promotion') == 'on'
    max_conditional = request.POST.get('max_conditional_subjects', 2)
    eval_term = request.POST.get('evaluation_term_sequence', 3)
    is_active = request.POST.get('is_active') == 'on'

    try:
        rule.promotion_mode = promotion_mode
        rule.minimum_passing_grade = min_pass
        rule.minimum_subjects_to_pass = min_subjects
        rule.minimum_overall_average = min_average
        rule.minimum_attendance_percentage = min_attendance
        rule.allow_conditional_promotion = allow_conditional
        rule.max_conditional_subjects = max_conditional
        rule.evaluation_term_sequence = eval_term
        rule.is_active = is_active
        rule.save()

        return JsonResponse({
            'success': True,
            'message': f"Promotion rule updated for {rule.from_grade_level.name} → {rule.to_grade_level.name}"
        })
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)})


@login_required
@require_POST
def promotion_rule_delete(request, rule_id):
    """Delete a promotion rule."""
    if request.user.role not in ['SUPER_ADMIN', 'SCHOOL_ADMIN']:
        return JsonResponse({'success': False, 'error': "Permission denied."}, status=403)

    school = request.user.school
    rule = get_object_or_404(PromotionRule, id=rule_id, school=school)

    try:
        rule.delete()
        return JsonResponse({'success': True, 'message': "Promotion rule deleted successfully."})
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)})
